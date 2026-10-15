from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

wb = Workbook()
ws = wb.active
ws.title = "Semester_Results"

ws.append(["Roll No", "Student Name", "Lab Marks (out of 50)", "Theory Marks (out of 50)", "Total Marks (out of 100)", "Result"])

students = [
    [1, "Aarav", 42, 38],
    [2, "Diya", 35, 30],
    [3, "Rohan", 15, 18],
    [4, "Meera", 45, 40],
    [5, "Kabir", 20, 25],
]

for student in students:
    ws.append(student)

last_row = ws.max_row

for row in range(2, last_row + 1):
    lab = ws.cell(row=row, column=3).value
    theory = ws.cell(row=row, column=4).value
    total = lab + theory
    ws.cell(row=row, column=5).value = total
    ws.cell(row=row, column=6).value = "Pass" if total >= 40 else "Fail"

for cell in ws[1]:
    cell.font = Font(bold=True)
    cell.fill = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")

total_sum = 0
for row in range(2, last_row + 1):
    total_sum += ws.cell(row=row, column=5).value

ws.cell(row=last_row + 1, column=2).value = "Class Average"
ws.cell(row=last_row + 1, column=5).value = total_sum / (last_row - 1)

wb.save("AIML_Batch_Grades.xlsx")